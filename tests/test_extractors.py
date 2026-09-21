"""Unit tests cho extractors.py."""

from pathlib import Path
import pytest

from src.extractors import extract_text, ExtractionError


def test_extract_text_from_txt(tmp_path: Path) -> None:
    txt_file = tmp_path / "note.txt"
    content = "Hệ thống thông tin quản lý\nNội dung bài học"
    txt_file.write_text(content, encoding="utf-8")

    extracted = extract_text(txt_file)
    assert "Hệ thống thông tin quản lý" in extracted
    assert "Nội dung bài học" in extracted


def test_extract_text_from_md(tmp_path: Path) -> None:
    md_file = tmp_path / "README.md"
    content = "# Tiêu đề\n- Gạch đầu dòng 1\n- Gạch đầu dòng 2"
    md_file.write_text(content, encoding="utf-8")

    extracted = extract_text(md_file)
    assert "# Tiêu đề" in extracted
    assert "- Gạch đầu dòng 1" in extracted


def test_extract_text_from_docx(tmp_path: Path) -> None:
    import docx

    docx_file = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Chương 1: Cơ sở dữ liệu", level=1)
    doc.add_paragraph("Mô hình dữ liệu quan hệ và chuẩn hóa dữ liệu.")
    doc.save(str(docx_file))

    extracted = extract_text(docx_file)
    assert "Chương 1: Cơ sở dữ liệu" in extracted
    assert "Mô hình dữ liệu quan hệ" in extracted


def test_extract_text_from_pptx(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    pptx_file = tmp_path / "slide.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Bài giảng Triết học Mác - Lênin"
    slide.placeholders[1].text = "Phần 1: Chủ nghĩa duy vật biện chứng"
    prs.save(str(pptx_file))

    extracted = extract_text(pptx_file)
    assert "[Slide 1]" in extracted
    assert "Bài giảng Triết học Mác - Lênin" in extracted
    assert "Chủ nghĩa duy vật biện chứng" in extracted


def test_extract_text_unsupported_extension(tmp_path: Path) -> None:
    unsupported = tmp_path / "binary.bin"
    unsupported.write_bytes(b"\x00\x01\x02\x03")

    with pytest.raises(ExtractionError) as exc_info:
        extract_text(unsupported)
    assert "chưa được hỗ trợ" in str(exc_info.value)


def test_extract_text_non_existent_file(tmp_path: Path) -> None:
    non_existent = tmp_path / "does_not_exist.txt"
    with pytest.raises(ExtractionError) as exc_info:
        extract_text(non_existent)
    assert "không tồn tại" in str(exc_info.value)
