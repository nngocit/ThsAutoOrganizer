"""Module trích xuất văn bản từ các định dạng tài liệu: PDF, DOCX, PPTX, TXT, MD.

Đảm bảo kiến trúc thống nhất, xử lý lỗi an toàn và sẵn sàng cho AI verification / OCR sau này.
"""

import logging
from pathlib import Path
from typing import Callable, Dict

logger = logging.getLogger("ThsAutoOrganizer.extractors")


class ExtractionError(Exception):
    """Lỗi xảy ra trong quá trình trích xuất nội dung văn bản."""
    pass


def extract_text_from_txt(path: Path) -> str:
    """Trích xuất text từ file plain text / markdown (.txt, .md)."""
    encodings = ["utf-8", "utf-8-sig", "cp1252", "latin-1"]
    for enc in encodings:
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
        except Exception as exc:
            raise ExtractionError(f"Không thể đọc file text '{path}': {exc}") from exc

    raise ExtractionError(f"Không thể giải mã file text '{path}' với các bộ mã phổ biến.")


def extract_text_from_pdf(path: Path) -> str:
    """Trích xuất text từ file PDF sử dụng pypdf."""
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ExtractionError(f"PDF được mã hóa / đặt mật khẩu: '{path}'") from exc

        pages_text = []
        for idx, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                pages_text.append(text.strip())

        return "\n\n".join(pages_text)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"Lỗi khi đọc file PDF '{path}': {exc}") from exc


def extract_text_from_docx(path: Path) -> str:
    """Trích xuất text từ file Word DOCX sử dụng python-docx."""
    try:
        import docx

        doc = docx.Document(str(path))
        text_elements = []

        # Đọc paragraphs
        for p in doc.paragraphs:
            if p.text.strip():
                text_elements.append(p.text.strip())

        # Đọc tables
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_elements.append(" | ".join(row_text))

        return "\n".join(text_elements)
    except Exception as exc:
        raise ExtractionError(f"Lỗi khi đọc file DOCX '{path}': {exc}") from exc


def extract_text_from_pptx(path: Path) -> str:
    """Trích xuất text từ file PowerPoint PPTX sử dụng python-pptx."""
    try:
        from pptx import Presentation

        prs = Presentation(str(path))
        slides_text = []

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_content = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_content.append(text)
            if slide_content:
                slides_text.append(f"[Slide {slide_idx}]\n" + "\n".join(slide_content))

        return "\n\n".join(slides_text)
    except Exception as exc:
        raise ExtractionError(f"Lỗi khi đọc file PPTX '{path}': {exc}") from exc


# Bảng đăng ký extractor theo extension
EXTRACTOR_REGISTRY: Dict[str, Callable[[Path], str]] = {
    ".txt": extract_text_from_txt,
    ".md": extract_text_from_txt,
    ".pdf": extract_text_from_pdf,
    ".docx": extract_text_from_docx,
    ".pptx": extract_text_from_pptx,
}


def extract_text(file_path: Path | str) -> str:
    """Interface thống nhất để trích xuất văn bản từ tài liệu.

    Args:
        file_path: Đường dẫn file cần trích xuất.

    Returns:
        Chuỗi văn bản đã trích xuất.

    Raises:
        ExtractionError: Khi file không được hỗ trợ hoặc gặp lỗi định dạng.
    """
    path = Path(file_path).resolve()

    if not path.is_file():
        raise ExtractionError(f"File không tồn tại: '{path}'")

    ext = path.suffix.lower()
    extractor = EXTRACTOR_REGISTRY.get(ext)

    if not extractor:
        supported = ", ".join(EXTRACTOR_REGISTRY.keys())
        raise ExtractionError(
            f"Định dạng file '{ext}' chưa được hỗ trợ trích xuất text. Các định dạng hỗ trợ: [{supported}]"
        )

    try:
        text = extractor(path)
        return text
    except ExtractionError:
        raise
    except Exception as exc:
        logger.error("Lỗi không mong đợi khi trích xuất text từ %s: %s", path, exc, exc_info=True)
        raise ExtractionError(f"Lỗi khi trích xuất text từ '{path}': {exc}") from exc
